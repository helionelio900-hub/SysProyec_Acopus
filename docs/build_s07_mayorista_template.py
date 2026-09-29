from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "plantilla_S07_mayorista.docx"
UPEU_LOGO = ROOT / "docs" / "presentacion" / "s06" / "assets" / "upeu.png"
SITRA_LOGO = ROOT / "docs" / "presentacion" / "s06" / "assets" / "logo-sitra-oro.jpg"

NAVY = "0B2742"
GOLD = "D99B13"
INK = "182B40"
MUTED = "627286"
PALE = "F3F6F9"
LINE = "CFD8E3"
WHITE = "FFFFFF"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(cell, color=LINE, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = tc_pr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        tc_pr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = b.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            b.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def font(run, size=10, color=INK, bold=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    return run


def para(p, before=0, after=5, line=1.12, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def text(doc, value="", size=10, color=INK, bold=False, before=0, after=5,
         align=None, italic=False):
    p = doc.add_paragraph()
    para(p, before, after, align=align)
    r = p.add_run(value)
    font(r, size, color, bold)
    r.italic = italic
    return p


def heading(doc, value, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    para(p, 5 if level == 1 else 2, 7 if level == 1 else 4)
    r = p.add_run(value)
    font(r, 18 if level == 1 else 12, NAVY, True)
    return p


def labeled_line(doc, label, width=83):
    p = doc.add_paragraph()
    para(p, 0, 7)
    font(p.add_run(label + "  "), 10, NAVY, True)
    font(p.add_run("_" * width), 10, MUTED)
    return p


def page_title(doc, number, title, intro):
    p = doc.add_paragraph()
    para(p, 0, 2)
    font(p.add_run(f"EVIDENCIA {number}  /  "), 9, GOLD, True)
    font(p.add_run("SESIÓN 07 · MÓDULO MAYORISTA"), 9, MUTED, True)
    heading(doc, title, 1)
    text(doc, intro, 10, MUTED, after=10)


def screenshot_box(doc, prompt, height_lines=8):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.35)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, PALE)
    borders(cell, LINE, "10")
    set_cell_margins(cell, top=100, bottom=100, start=150, end=150)
    p = cell.paragraphs[0]
    para(p, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("ESPACIO PARA PEGAR LA CAPTURA"), 10, NAVY, True)
    p = cell.add_paragraph()
    para(p, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run(prompt), 9, MUTED)
    for _ in range(height_lines):
        p = cell.add_paragraph()
        para(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        font(p.add_run(" "), 9, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def caption_fields(doc):
    labeled_line(doc, "Qué se ve y qué demuestra:", 68)
    labeled_line(doc, "Fecha y hora de la captura:", 66)


def checkbox(doc, label):
    p = doc.add_paragraph()
    para(p, 0, 4)
    font(p.add_run("☐  "), 11, GOLD, True)
    font(p.add_run(label), 10, INK)


def page_break(doc):
    doc.add_page_break()


def field_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    font(run, 8, MUTED)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for name in ("Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "Aptos Display"
        st.font.color.rgb = RGBColor.from_string(NAVY)
        st.font.bold = True
    doc.styles["Title"].font.name = "Aptos Display"
    doc.styles["Title"].font.size = Pt(27)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].font.color.rgb = RGBColor.from_string("000000")

    footer = sec.footer.paragraphs[0]
    para(footer, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
    font(footer.add_run("SITRA-ORO  ·  S07  ·  Evidencia individual                         Página "), 8, MUTED)
    field_page_number(footer)

    # Cover
    logos = doc.add_table(rows=1, cols=2)
    logos.autofit = False
    logos.columns[0].width = Inches(3.25)
    logos.columns[1].width = Inches(3.25)
    for c in logos.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = logos.cell(0, 0).paragraphs[0]
    para(p, 0, 0, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.add_run().add_picture(str(UPEU_LOGO), width=Inches(2.65))
    p = logos.cell(0, 1).paragraphs[0]
    para(p, 0, 0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p.add_run().add_picture(str(SITRA_LOGO), width=Inches(1.48))
    text(doc, "UNIVERSIDAD PERUANA UNIÓN", 10, NAVY, True, before=44, after=3,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    text(doc, "Escuela Profesional de Ingeniería de Sistemas", 10, MUTED, after=28,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph(style="Title")
    para(p, 0, 7, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("Evidencia individual"), 27, "000000", True)
    p = doc.add_paragraph()
    para(p, 0, 4, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("Sesión 07 · Arquitectura SPA y CRUD"), 16, NAVY, True)
    p = doc.add_paragraph()
    para(p, 0, 23, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    font(p.add_run("Aplicación al módulo Mayorista de SITRA-ORO"), 12, MUTED)
    cover = doc.add_table(rows=3, cols=2)
    cover.autofit = False
    cover.columns[0].width = Inches(1.55)
    cover.columns[1].width = Inches(4.75)
    for row, label in zip(cover.rows, ("Estudiante", "Equipo", "Fecha de entrega")):
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(4.75)
        for c in row.cells:
            borders(c, LINE, "7")
            set_cell_margins(c)
        shade(row.cells[0], "EAF0F5")
        para(row.cells[0].paragraphs[0], 0, 0)
        font(row.cells[0].paragraphs[0].add_run(label), 9, NAVY, True)
        para(row.cells[1].paragraphs[0], 0, 0)
        font(row.cells[1].paragraphs[0].add_run("Completar aquí"), 9, MUTED)
    text(doc, "Lenguaje de Programación II", 10, NAVY, True, before=38,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    text(doc, "SITRA-ORO · Módulo Mayorista", 9, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Scope and checklist
    page_break(doc)
    heading(doc, "Datos y alcance de mi trabajo", 1)
    text(doc, "Completo esta ficha con lo que implementé y probé personalmente. No marco como terminado algo que aún no comprobé.", 10, MUTED, after=12)
    labeled_line(doc, "Nombre completo:", 74)
    labeled_line(doc, "Equipo:", 82)
    labeled_line(doc, "Aporte realizado:", 73)
    labeled_line(doc, "Enlace del repositorio:", 67)
    heading(doc, "CRUD que voy a demostrar", 2)
    labeled_line(doc, "Nombre de la funcionalidad:", 59)
    labeled_line(doc, "Entidad o tabla del backend:", 59)
    labeled_line(doc, "Servicio frontend:", 73)
    text(doc, "Antes de capturar, confirmaré que el CRUD elegido cumple la condición de la actividad: no requiere seleccionar previamente un registro relacionado. Si tengo dudas, validaré el alcance con el docente.", 10, INK, after=7)
    checkbox(doc, "Probé la función con el backend real, no con vista de demostración.")
    checkbox(doc, "Tengo disponibles crear, listar, editar y eliminar para la entidad elegida.")
    checkbox(doc, "Puedo explicar qué dato se guarda y qué validación aplica.")
    heading(doc, "Orden de la evidencia", 2)
    text(doc, "Las capturas de las páginas siguientes deben tomarse durante tu trabajo real. Antes de cada captura, deja visibles el reloj de Windows y tu usuario o perfil en VS Code o en el navegador; no recortes esos elementos. Debajo escribe qué se observa con tus propias palabras.", 10, INK)

    # Block 1
    page_break(doc)
    page_title(doc, "01", "Proyecto y arquitectura", "Muestra la estructura real del frontend y señala dónde viven los elementos comunes y la funcionalidad Mayorista.")
    checkbox(doc, "Se distinguen core, shared y features.")
    checkbox(doc, "Se identifica la carpeta de la funcionalidad Mayorista.")
    screenshot_box(doc, "Explorador de VS Code con las carpetas del proyecto visibles.", 10)
    caption_fields(doc)
    text(doc, "Explicación", 10, NAVY, True, before=3, after=2)
    labeled_line(doc, "¿Cómo organicé el módulo Mayorista?", 53)

    # Block 2
    page_break(doc)
    page_title(doc, "02", "Layout y navegación", "La aplicación debe abrir una página de inicio real en / y conservar el encabezado, menú y navegación al cambiar de ruta.")
    checkbox(doc, "La ruta / muestra el inicio; no redirige automáticamente a otra pantalla.")
    checkbox(doc, "Se ven encabezado, menú y área de contenido.")
    checkbox(doc, "Puedo navegar entre al menos dos rutas.")
    screenshot_box(doc, "Captura de la aplicación en ejecución, con el reloj y el perfil visibles.", 11)
    caption_fields(doc)
    labeled_line(doc, "Rutas que recorrí:", 70)

    # Block 3
    page_break(doc)
    page_title(doc, "03", "Servicio HTTP e interceptor", "Registra la separación entre el servicio de la funcionalidad, la URL base y el identificador que acompaña las solicitudes.")
    text(doc, "Captura A · Código", 10, NAVY, True, after=3)
    screenshot_box(doc, "Servicio de funcionalidad y ApiService en VS Code.", 5)
    text(doc, "Captura B · Solicitud real", 10, NAVY, True, before=4, after=3)
    screenshot_box(doc, "Network → Request Headers; debe verse X-Trace-ID y una respuesta exitosa del backend.", 5)
    labeled_line(doc, "Qué petición real probé y qué respuesta recibí:", 47)

    # Block 4a
    page_break(doc)
    page_title(doc, "04", "CRUD contra el backend real", "Documenta los cuatro casos con datos de prueba que puedas usar sin exponer información sensible.")
    text(doc, "Crear", 12, NAVY, True, after=3)
    screenshot_box(doc, "Formulario completado y confirmación de creación exitosa.", 6)
    labeled_line(doc, "Qué registro creé:", 70)
    text(doc, "Listar", 12, NAVY, True, before=8, after=3)
    screenshot_box(doc, "Registro recién creado visible en la lista obtenida del backend.", 6)
    labeled_line(doc, "Cómo comprobé que vino del backend:", 53)

    # Block 4b
    page_break(doc)
    page_title(doc, "04", "CRUD contra el backend real", "Completa esta segunda página para mostrar que los cambios y la eliminación también se reflejan sin recargar manualmente.")
    text(doc, "Editar", 12, NAVY, True, after=3)
    screenshot_box(doc, "Registro modificado y confirmación de actualización exitosa.", 6)
    labeled_line(doc, "Qué dato cambié:", 71)
    text(doc, "Eliminar", 12, NAVY, True, before=8, after=3)
    screenshot_box(doc, "Confirmación de eliminación y lista actualizada.", 6)
    labeled_line(doc, "Cómo confirmé que ya no aparece:", 57)

    # Finding and reflection
    page_break(doc)
    heading(doc, "Hallazgo técnico real", 1)
    text(doc, "Describe un problema que sí encontraste al trabajar. Explica qué ocurrió, qué lo causaba y cómo comprobaste la solución. No inventes un error para llenar este espacio.", 10, MUTED, after=10)
    labeled_line(doc, "Qué ocurrió:", 75)
    labeled_line(doc, "Causa que encontré:", 67)
    labeled_line(doc, "Cómo lo resolví o qué falta resolver:", 50)
    screenshot_box(doc, "Evidencia del error o del resultado después de corregirlo.", 5)
    heading(doc, "Reflexión técnica", 1)
    text(doc, "¿Por qué separar el servicio de la funcionalidad del componente facilita un cambio futuro en la URL o en la forma de consumir el backend? Responde entre 5 y 8 líneas y menciona tu servicio real.", 10, INK, after=4)
    for _ in range(6):
        labeled_line(doc, "", 84)

    # Oral preparation prompts based on the assignment, with domain terms
    page_break(doc)
    heading(doc, "Preguntas para explicar mi trabajo", 1)
    text(doc, "Úsalas para preparar tu sustentación. Respóndelas con decisiones y nombres que existan en tu proyecto.", 10, MUTED, after=9)
    questions = [
        "¿Por qué el layout vive separado de las pantallas de cada función?",
        "¿Qué ventaja tiene organizar por core, shared y features frente a juntar todo por tipo?",
        "¿Por qué el componente usa un servicio de funcionalidad y no llama directamente a HttpClient?",
        "¿Por qué el CRUD que elegí cumple —o no cumple todavía— la condición de independencia?",
        "¿Por qué ApiService no conoce las reglas del Mayorista y qué otro servicio lo puede reutilizar?",
        "¿Por qué el interceptor agrega X-Trace-ID en un solo lugar? ¿Qué permite rastrear?",
        "¿Qué validaciones del backend tuve que respetar en el formulario?",
    ]
    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        para(p, 1, 1)
        font(p.add_run(f"{i}. {q}"), 9.5, NAVY, True)
        labeled_line(doc, "", 82)

    # Feedback must remain the final page as specified by the assignment
    page_break(doc)
    heading(doc, "Anexo de feedback de la sesión", 1)
    text(doc, "Responde con honestidad. Esta página debe quedar al final del PDF entregado.", 10, MUTED, after=8)
    feedback = [
        "¿Cuál es el aprendizaje más importante que te llevas de la clase de hoy?",
        "¿Qué punto de la clase te resultó más confuso o te dejó con dudas?",
        "¿Tienes alguna pregunta que te gustaría que se responda la siguiente clase?",
    ]
    for i, q in enumerate(feedback, 1):
        p = doc.add_paragraph()
        para(p, 2, 1)
        font(p.add_run(f"{i}. {q}"), 9.5, NAVY, True)
        labeled_line(doc, "", 81)
    p = doc.add_paragraph()
    para(p, 3, 2)
    font(p.add_run("4. Mi comprensión de la clase:"), 9.5, NAVY, True)
    for opt in ("Entendido: podría explicarlo.", "Más o menos: entendí la idea, aún tengo dudas.", "Necesito ayuda: todavía me siento perdido/a."):
        checkbox(doc, opt)
    for i, q in ((5, "¿Cómo pueden ayudarte a comprender mejor el tema?"),
                 (6, "¿Cómo evaluarías tu participación y esfuerzo?")):
        p = doc.add_paragraph()
        para(p, 2, 1)
        font(p.add_run(f"{i}. {q}"), 9.5, NAVY, True)
        labeled_line(doc, "", 81)
    for opt in ("Muy comprometido/a: me esforcé al máximo.",
                "Comprometido/a: podría haberme esforzado un poco más.",
                "Poco comprometido/a: hoy no di mi mejor esfuerzo."):
        checkbox(doc, opt)
    p = doc.add_paragraph()
    para(p, 2, 1)
    font(p.add_run("7. Mi satisfacción con la clase (del 1 al 10):"), 9.5, NAVY, True)
    text(doc, "1   2   3   4   5   6   7   8   9   10        Marca una opción", 10, INK, after=0)

    doc.core_properties.title = "Plantilla de evidencia individual S07 del módulo Mayorista"
    doc.core_properties.subject = "Arquitectura SPA y CRUD para SITRA-ORO"
    doc.core_properties.author = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
