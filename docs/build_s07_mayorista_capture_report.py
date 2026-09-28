from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "S07_Mayorista_Plantilla_Capturas.docx"
ASSETS = ROOT / "docs" / "presentacion" / "s06" / "assets"
BLUE = "203F87"
NAVY = "102849"
GOLD = "D7A11E"
INK = "1E2B3C"
MUTED = "607087"
PALE = "F5F8FC"
LINE = "9BAFCC"


def set_font(run, size=9.5, color=INK, bold=False, italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def format_p(p, before=0, after=4, line=1.08, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_text(doc, value, size=9.5, color=INK, bold=False, before=0, after=4,
             align=None, italic=False):
    p = doc.add_paragraph()
    format_p(p, before, after, align=align)
    set_font(p.add_run(value), size, color, bold, italic)
    return p


def section_title(doc, title, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        format_p(p, 0, 1)
        set_font(p.add_run(kicker.upper()), 8, GOLD, True)
    p = doc.add_paragraph(style="Heading 1")
    format_p(p, 0, 5)
    set_font(p.add_run(title), 17, NAVY, True)
    return p


def line(doc, label, underscores=72):
    p = doc.add_paragraph()
    format_p(p, 0, 5)
    set_font(p.add_run(label + "  "), 9, NAVY, True)
    set_font(p.add_run("_" * underscores), 9, MUTED)


def add_checkbox(doc, label):
    p = doc.add_paragraph()
    format_p(p, 0, 2)
    set_font(p.add_run("☐  "), 10, GOLD, True)
    set_font(p.add_run(label), 9, INK)


def screenshot_slot(doc, number, title, instruction, lines=7, explanation=""):
    p = doc.add_paragraph()
    format_p(p, 5, 1)
    set_font(p.add_run(f"CAPTURA {number}. {title}"), 9.5, BLUE, True)
    add_text(doc, instruction, 8.5, MUTED, italic=True, after=3)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.45)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), PALE)
    tc_pr.append(shade)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement("w:" + edge)
        tag.set(qn("w:val"), "dashed")
        tag.set(qn("w:sz"), "9")
        tag.set(qn("w:color"), BLUE)
        borders.append(tag)
    tc_pr.append(borders)
    margins = OxmlElement("w:tcMar")
    for edge, amount in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), amount)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)
    p = cell.paragraphs[0]
    format_p(p, 0, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("[ PEGA AQUÍ LA CAPTURA ]"), 10, BLUE, True)
    p = cell.add_paragraph()
    format_p(p, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("Ventana completa · reloj y usuario/perfil visibles · sin recortar"), 8, MUTED)
    for _ in range(lines):
        p = cell.add_paragraph()
        format_p(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(" "), 8)
    if explanation:
        p = doc.add_paragraph()
        format_p(p, 2, 4)
        set_font(p.add_run("Explicación. "), 8.5, NAVY, True)
        set_font(p.add_run(explanation), 8.5, INK)


def page_number(p):
    r = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, end])
    set_font(r, 8, MUTED)


def new_page(doc):
    doc.add_page_break()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
    title_style = doc.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")

    header = sec.header.paragraphs[0]
    format_p(header, 0, 0)
    set_font(header.add_run("SITRA-ORO · Evidencia S07"), 8, MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = sec.footer.paragraphs[0]
    format_p(footer, 0, 0)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Faijo Calisaya Helio Paul · Equipo 05                         Página "), 8, MUTED)
    page_number(footer)

    # Portada inspirada en el PDF de referencia, con los tres logos facilitados.
    logos = doc.add_table(rows=1, cols=3)
    logos.autofit = False
    widths = (Inches(2.1), Inches(1.35), Inches(1.0))
    paths = (ASSETS / "upeu.png", ASSETS / "quilate-bullion-emblema-v2.png", ASSETS / "ingenieria-sistemas.png")
    sizes = (Inches(1.55), Inches(0.83), Inches(0.68))
    for i, (cell, path, width) in enumerate(zip(logos.rows[0].cells, paths, widths)):
        cell.width = width
        p = cell.paragraphs[0]
        format_p(p, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(str(path), width=sizes[i])
    p = doc.add_paragraph()
    format_p(p, 14, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("UNIVERSIDAD PERUANA UNIÓN"), 10, NAVY, True)
    add_text(doc, "Escuela Profesional de Ingeniería de Sistemas", 9, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    p = doc.add_paragraph(style="Title")
    format_p(p, 0, 5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("Informe de evidencia de aprendizaje"), 24, NAVY, True)
    add_text(doc, "S07 · Creación y arquitectura de la SPA", 14, BLUE, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_text(doc, "SITRA-ORO · Aplicación al módulo Mayorista", 11, INK, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    facts = doc.add_table(rows=5, cols=2)
    facts.autofit = False
    facts.columns[0].width = Inches(1.25)
    facts.columns[1].width = Inches(5.1)
    fields = [
        ("Nombre", "Faijo Calisaya Helio Paul"),
        ("Equipo", "Equipo 05"),
        ("Sesión", "S07 · Creación y Arquitectura de la SPA"),
        ("Aporte", "Frontend del módulo Mayorista: arquitectura SPA, navegación y conexión HTTP."),
        ("GitHub", "https://github.com/helionelio900-hub/SysProyec_Acopus"),
    ]
    for row, (label, value) in zip(facts.rows, fields):
        for i, c in enumerate(row.cells):
            tc_pr = c._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), BLUE if i == 0 else "FFFFFF")
            tc_pr.append(shd)
            b = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "5")
                e.set(qn("w:color"), "C7D3E3")
                b.append(e)
            tc_pr.append(b)
        p = row.cells[0].paragraphs[0]
        format_p(p, 0, 0)
        set_font(p.add_run(label), 8.5, "FFFFFF", True)
        p = row.cells[1].paragraphs[0]
        format_p(p, 0, 0)
        set_font(p.add_run(value), 8.5, INK, value == "Equipo 05")
    section_title(doc, "Propósito", "2 · Aplicación de la sesión")
    add_text(doc, "Documentar la arquitectura SPA del frontend SITRA-ORO y comprobar la navegación y las operaciones del módulo Mayorista conectado al backend.", 10, INK, after=7)
    add_text(doc, "Las capturas son la evidencia visual pendiente. Inserta capturas reales con la ventana completa, la fecha y hora, y tu usuario o perfil visibles.", 9, MUTED, before=2, after=5)

    # 1. Arquitectura
    new_page(doc)
    section_title(doc, "Proyecto y arquitectura", "3.1 · Evidencia técnica")
    add_text(doc, "El frontend de SITRA-ORO está organizado por responsabilidades comunes y funcionalidades. En esta evidencia se debe mostrar la estructura real del proyecto, incluyendo core, shared y features, y la ubicación de las pantallas del módulo Mayorista.", 9.5, after=6)
    add_text(doc, "Ruta de referencia: lp2/sitra-oro-frontend/src/app", 8.5, MUTED, italic=True)
    screenshot_slot(doc, 1, "Estructura de carpetas del proyecto", "En VS Code, muestra core, shared, features y la carpeta Mayorista. Deja visible el reloj y tu perfil.", 13,
                   "La estructura separa los elementos compartidos de la aplicación de las pantallas y servicios propios del módulo Mayorista.")

    # 2. Layout and routes
    new_page(doc)
    section_title(doc, "Layout y navegación", "3.2 · Evidencia técnica")
    add_text(doc, "La actividad pide una página de inicio real en / y al menos dos rutas navegables. Las capturas deben mostrar que el encabezado y menú se mantienen al cambiar entre páginas.", 9.5, after=4)
    screenshot_slot(doc, 2, "Página de inicio en la ruta raíz", "Abre http://localhost:4200/; muestra el inicio real, el encabezado y el menú.", 6,
                   "La página raíz presenta el inicio del proyecto sin enviar al usuario automáticamente a otra sección.")
    screenshot_slot(doc, 3, "Navegación entre dos rutas", "Navega a una ruta Mayorista y conserva el layout común visible.", 6,
                   "La navegación cambia el contenido de la ruta seleccionada y mantiene el encabezado y el menú comunes.")

    # 3. HTTP / interceptor
    new_page(doc)
    section_title(doc, "Servicio HTTP e interceptor", "3.3 · Evidencia técnica")
    add_text(doc, "El servicio de funcionalidad debe concentrar las llamadas del módulo, usar ApiService para formar las URL y dejar que el interceptor agregue X-Trace-ID a cada petición saliente.", 9.5, after=3)
    screenshot_slot(doc, 4, "Código del servicio y el interceptor", "Muestra en VS Code el servicio Mayorista y traceIdInterceptor con crypto.randomUUID().", 6,
                   "OperacionesService concentra las llamadas del módulo y el interceptor añade un identificador de seguimiento a cada solicitud.")
    screenshot_slot(doc, 5, "Petición real en Network", "Muestra Request Headers con X-Trace-ID y el resultado real de la petición al backend.", 6,
                   "La pestaña Network permite comprobar la cabecera X-Trace-ID y el resultado devuelto por el backend para esa petición.")

    # 4a. CRUD list/create
    new_page(doc)
    section_title(doc, "CRUD contra el backend real", "3.4 · Evidencia técnica")
    add_text(doc, "Inserta capturas de las operaciones que ejecutaste contra tu backend real. No uses la vista de demostración ni afirmes una operación que todavía no probaste.", 9.5, after=4)
    screenshot_slot(doc, 6, "Listar", "Muestra los registros cargados desde el backend y, si es posible, la petición exitosa en Network.", 6,
                   "El listado debe reflejar los registros recibidos del servidor, no datos escritos únicamente en la pantalla.")
    screenshot_slot(doc, 7, "Crear", "Muestra el envío del formulario y la confirmación del registro persistido.", 6,
                   "La creación se demuestra con una respuesta exitosa del servidor y el nuevo registro visible en la lista.")

    # 4b. CRUD update/delete
    new_page(doc)
    section_title(doc, "CRUD contra el backend real", "3.4 · Evidencia técnica")
    screenshot_slot(doc, 8, "Editar", "Muestra el registro actualizado y la respuesta del backend.", 6,
                   "La edición debe persistir los cambios en el servidor y mostrarlos en la interfaz sin una recarga manual.")
    screenshot_slot(doc, 9, "Eliminar", "Muestra la confirmación y la lista actualizada sin el registro eliminado.", 6,
                   "La eliminación debe confirmarse en el servidor y retirar el registro de la lista visible.")

    # Finding / reflection
    new_page(doc)
    section_title(doc, "Error o hallazgo técnico", "4 · Análisis")
    add_text(doc, "En la pantalla de recepción apareció un mensaje genérico al fallar la consulta de registros.", 9.5, after=4)
    screenshot_slot(doc, 10, "Hallazgo en la carga de recepciones", "Muestra el aviso real o la respuesta de Network que permitió reconocer el problema.", 6,
                   "El componente agrupaba varios errores bajo un mismo aviso, por lo que no orientaba sobre la causa. Se ajustó para distinguir sesión vencida, falta de permiso, ruta no disponible y otros fallos de conexión.")
    section_title(doc, "Reflexión técnica breve", "5 · Cierre")
    add_text(doc, "¿Por qué separar el servicio de la funcionalidad del componente facilita un cambio futuro en la URL o en la forma de consumir el backend? Responde en 5 a 8 líneas y menciona un servicio real de tu proyecto.", 9, after=3)
    add_text(doc, "Separar OperacionesService del componente permite que el componente se concentre en la pantalla y que el servicio gestione las peticiones HTTP. Si cambia una ruta o la forma de consumir el backend, el ajuste se puede centralizar en el servicio, sin repetir cambios en cada pantalla. También mantiene tipadas las respuestas y facilita probar por separado la vista y la comunicación con el servidor.", 9, INK, after=4)

    # Required feedback is the final page, with the answers from the user's supplied reference.
    new_page(doc)
    p = doc.add_paragraph()
    format_p(p, 0, 1, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("UNIVERSIDAD PERUANA UNIÓN"), 9, NAVY, True)
    add_text(doc, "Escuela Profesional de Ingeniería de Sistemas", 8.5, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    section_title(doc, "Anexo · Feedback de la sesión", "Última página del informe")
    items = [
        "¿Cuál es el aprendizaje más importante que te llevas de la clase de hoy?",
        "¿Qué punto de la clase te resultó más confuso o te dejó con dudas?",
        "¿Qué pregunta te gustaría que se responda en la siguiente clase?",
    ]
    answers = [
        "Aprendí a organizar una aplicación Angular en core, shared y features, y a conectar las pantallas con servicios e interceptores HTTP.",
        "Me costó comprender al inicio cómo las rutas hijas se muestran dentro del layout sin reemplazar el encabezado y el menú.",
        "¿Cómo se renueva un token JWT cuando varias solicitudes reciben un error 401 al mismo tiempo?",
    ]
    for i, (prompt, answer) in enumerate(zip(items, answers), 1):
        p = doc.add_paragraph()
        format_p(p, 3, 0)
        set_font(p.add_run(f"{i}. {prompt}"), 9, NAVY, True)
        add_text(doc, answer, 8.7, INK, after=2)
    p = doc.add_paragraph()
    format_p(p, 3, 1)
    set_font(p.add_run("4. Mi comprensión de la clase:"), 9, NAVY, True)
    for answer in ("[X] Entendido; podría explicarlo.", "[ ] Más o menos; tengo algunas dudas.", "[ ] Necesito ayuda con el tema."):
        add_text(doc, answer, 8.7, INK, after=1)
    for i, prompt in ((5, "¿Cómo pueden ayudarte a comprender mejor el tema?"), (6, "¿Cómo evaluarías tu participación y esfuerzo?")):
        p = doc.add_paragraph()
        format_p(p, 3, 0)
        set_font(p.add_run(f"{i}. {prompt}"), 9, NAVY, True)
        if i == 5:
            add_text(doc, "Con ejemplos prácticos de Signals y del manejo de estados de carga.", 8.7, INK, after=2)
    for answer in ("[X] Muy comprometido/a.", "[ ] Comprometido/a.", "[ ] Poco comprometido/a."):
        add_text(doc, answer, 8.7, INK, after=1)
    p = doc.add_paragraph()
    format_p(p, 3, 1)
    set_font(p.add_run("7. Mi satisfacción con la clase (del 1 al 10):"), 9, NAVY, True)
    add_text(doc, "1     2     3     4     5     6     7     8     9     [X] 10", 9)

    doc.core_properties.title = "Informe de evidencia S07 del módulo Mayorista"
    doc.core_properties.subject = "Plantilla para insertar capturas de evidencia"
    doc.core_properties.author = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
