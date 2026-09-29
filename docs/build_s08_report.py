"""Rebuild the S08 report from the retained S07 Mayorista document system."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
from hashlib import sha256

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
REF = DOCS / 'S07_Mayorista_S07-FaijoCalisayaHelioPaul.docx'
OUT = DOCS / 'S08_LP2_Equipo05_FaijoCalisayaHelioPaul.docx'
MEDIA = DOCS / 'capturas' / 'S07_reference_media'
ASSETS = DOCS / 'presentacion' / 's06' / 'assets'
EXPECTED_REF_SHA = 'CCA19012720DA4E6A4BA4FFC1A72B8555D84331E9F893BFAE27E848758EACF69'

if sha256(REF.read_bytes()).hexdigest().upper() != EXPECTED_REF_SHA:
    raise RuntimeError('La referencia S07 cambió desde la revisión; no se modificará ni usará como plantilla sin nueva revisión.')

NAVY = '102849'
BODY = '1E2B3C'
GOLD = 'C88B00'
GOLD_PALE = 'FFF8E9'
BLUE = '315F91'
MUTED = '596A7B'
PALE = 'F4F7FA'
RED = 'A43B30'
WHITE = 'FFFFFF'

doc = Document(REF)  # Working copy in memory; the retained S07 file is untouched.
doc._body.clear_content()
sec = doc.sections[0]
sec.top_margin = Inches(.62)
sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.72)
sec.right_margin = Inches(.72)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(BODY)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for stylename, size, color in [('Title', 25, NAVY), ('Heading 1', 14, NAVY), ('Heading 2', 13, NAVY)]:
    style = styles[stylename]
    style.font.name = 'Aptos Display' if stylename in ('Title', 'Heading 1') else 'Aptos'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(5)

def run(p, text, *, bold=False, italic=False, color=BODY, size=9.5, font='Aptos'):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r

def p(text='', *, style='Normal', align=None, before=0, after=5, keep=False):
    q = doc.add_paragraph(style=style)
    if text:
        q.add_run(text)
    if align is not None:
        q.alignment = align
    q.paragraph_format.space_before = Pt(before)
    q.paragraph_format.space_after = Pt(after)
    q.paragraph_format.keep_together = keep
    return q

def kicker(text):
    q = p(align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=3, keep=True)
    run(q, text.upper(), bold=True, color=GOLD, size=8.5)
    return q

def heading(text, level=1):
    q = doc.add_paragraph(text, style=f'Heading {level}')
    q.paragraph_format.keep_with_next = True
    return q

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def cell_margins(cell, top=95, start=115, bottom=95, end=115):
    tc_pr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = OxmlElement(f'w:{edge}'); node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa'); margins.append(node)
    tc_pr.append(margins)

def table_borders(table, color='B9C7D5', size='5'):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color); borders.append(el)
    tbl_pr.append(borders)

def set_repeat_header(row):
    tr_pr=row._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); tr_pr.append(h)

def set_row_height(row, inches, exact=False):
    tr_pr=row._tr.get_or_add_trPr(); h=OxmlElement('w:trHeight'); h.set(qn('w:val'),str(round(inches*1440))); h.set(qn('w:hRule'),'exact' if exact else 'atLeast'); tr_pr.append(h)

def image_paragraph(path, width, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2):
    q=p(align=align,before=before,after=after,keep=True); q.add_run().add_picture(str(path),width=Inches(width)); return q

def image_box(number, title, instruction, height=2.35):
    # One editable framed slot, visually matching the S07 screenshot blocks.
    titlep=p(after=2,keep=True); run(titlep,f'CAPTURA {number}. {title}',bold=True,color=NAVY,size=9.2)
    inst=p(after=4,keep=True); run(inst,instruction,italic=True,color=MUTED,size=8.1)
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(6.9); c=t.cell(0,0); c.width=Inches(6.9); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(c,PALE); cell_margins(c,120,150,120,150); table_borders(t)
    set_row_height(t.rows[0],height,False)
    q=c.paragraphs[0]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(4)
    run(q,f'[ PEGA AQUÍ LA CAPTURA {number} ]',bold=True,color=BLUE,size=10)
    q=c.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(0)
    run(q,'Ventana completa · reloj del sistema y usuario/perfil visibles · sin recortar',italic=True,color=MUTED,size=7.7)
    cp=p(after=7,before=3,keep=True); run(cp,'Explicación. ',bold=True,color=NAVY,size=8.7); run(cp,instruction,color=BODY,size=8.7)
    return t

def data_table(headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    table_borders(t,'C5D0DC','5'); set_repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); cell_margins(c,85,100,85,100); run(c.paragraphs[0],h,bold=True,color=WHITE,size=8.4)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            c=cells[i]; cell_margins(c,75,100,75,100)
            if ri%2==1: shade(c,'F5F7FA')
            c.text=''; run(c.paragraphs[0],value,color=BODY,size=8.6)
    return t

def callout(title, body, *, fill=GOLD_PALE, title_color=GOLD):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,fill); cell_margins(c,105,140,105,140); table_borders(t,color='E8D9B7',size='5')
    q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(2); run(q,title,bold=True,color=title_color,size=8.8)
    q=c.add_paragraph(); q.paragraph_format.space_after=Pt(0); run(q,body,color=BODY,size=8.7)
    p('',after=2)

def set_page_footer(section, page_number):
    section.footer.is_linked_to_previous=False
    footer=section.footer.paragraphs[0]
    footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run(footer,f'Faijo Calisaya Helio Paul · Equipo 05                              Página {page_number}',color=MUTED,size=8)

page_number=1
def page_break():
    global page_number
    page_number += 1
    section=doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin=sec.top_margin; section.bottom_margin=sec.bottom_margin
    section.left_margin=sec.left_margin; section.right_margin=sec.right_margin
    set_page_footer(section,page_number)

# Reuse the reference's internal-page chrome. Each page has its own footer
# instance with a fixed visible page number, avoiding stale Word field caches.
header=sec.header.paragraphs[0]; header.clear(); header.alignment=WD_ALIGN_PARAGRAPH.LEFT
run(header,'SITRA-ORO · Evidencia S08',color=BLUE,size=8)
set_page_footer(sec,1)

# PAGE 1 — cover and purpose, based on S07's institutional cover.
logos=doc.add_table(rows=1,cols=2); logos.alignment=WD_TABLE_ALIGNMENT.CENTER; logos.autofit=False
for col in logos.columns: col.width=Inches(3.37)
for c in logos.rows[0].cells: c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
logo_specs=[(ASSETS/'upeu.png',1.55,WD_ALIGN_PARAGRAPH.LEFT),(ASSETS/'ingenieria-sistemas.png',.68,WD_ALIGN_PARAGRAPH.RIGHT)]
for i,(path,width,alignment) in enumerate(logo_specs):
    q=logos.cell(0,i).paragraphs[0]; q.alignment=alignment; q.paragraph_format.space_after=Pt(0)
    q.add_run().add_picture(str(path),width=Inches(width))
image_paragraph(MEDIA/'image1.png',1.65,before=5,after=2)
q=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=0); run(q,'UNIVERSIDAD PERUANA UNIÓN',bold=True,color=NAVY,size=9.5)
q=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=4); run(q,'Escuela Profesional de Ingeniería de Sistemas',color=BLUE,size=9)
q=doc.add_paragraph(style='Title'); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(4)
q.add_run('Informe de evidencia de aprendizaje')
q=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=4); run(q,'S08 · CRUD de tablas dependientes',bold=True,color=NAVY,size=15)
q=p(align=WD_ALIGN_PARAGRAPH.CENTER,after=8); run(q,'SITRA-ORO · Recepción Mayorista y Centro de Acopio',bold=True,color=BODY,size=10.5)

student=doc.add_table(rows=5,cols=2); student.alignment=WD_TABLE_ALIGNMENT.CENTER; student.autofit=False
student.columns[0].width=Inches(1.5); student.columns[1].width=Inches(5.25)
rows=[('Nombre','Faijo Calisaya Helio Paul'),('Equipo','Equipo 05'),('Sesión','S08 - CRUD de Tablas Dependientes'),('Aporte','Vínculo entre recepciones mayoristas y centros de acopio; validación y migración de datos.'),('GitHub','https://github.com/helionelio900-hub/SysProyec_Acopus')]
table_borders(student,'B7C9DE','5')
for i,(lab,val) in enumerate(rows):
    a,b=student.rows[i].cells; shade(a,NAVY); shade(b,'FFFFFF'); cell_margins(a,65,90,65,90); cell_margins(b,65,100,65,100)
    run(a.paragraphs[0],lab,bold=True,color=WHITE,size=8.3); run(b.paragraphs[0],val,color=BLUE if lab in ('Nombre','GitHub') else BODY,size=8.3)
kicker('2 · Aplicación de la sesión')
heading('Propósito',1)
p('Documentar un CRUD dependiente del dominio SITRA-ORO, explicar la relación entre Recepción Mayorista y Centro de Acopio, y mostrar cómo el identificador del padre conecta el formulario, el backend y la base de datos.',after=4)
p('Las capturas se insertarán manualmente. Cada una debe mostrar la ventana completa, el reloj de Windows con fecha y hora y el usuario o perfil del estudiante, sin recortes.',after=0)

# PAGE 2 — relation, models and services.
page_break(); kicker('3.1 · Evidencia técnica'); heading('Modelos y servicios HTTP',1)
p('La tabla dependiente elegida es RecepcionMayorista. Cada recepción pertenece a un Centro de Acopio, que funciona como tabla padre. El request envía idCentroAcopio como número; el response devuelve el identificador y una etiqueta legible del centro para mostrarla en la interfaz.')
data_table(['Elemento','Estructura en SITRA-ORO','Propósito'],[
('Padre','CentroAcopio: idCentroAcopio, nombre, zona, activo','Provee opciones válidas para la recepción.'),
('Request','CrearRecepcionMayorista: fecha, idCentroAcopio, rojo, verde, descuento, total','Representa los datos que recibe POST/PUT.'),
('Response','RecepcionMayorista: idRecepcion, fecha, idCentroAcopio, nombreAcopiador, rojo, verde, descuento, total','Agrega la identidad del registro y datos relacionados para consulta.'),
('Relación JPA/SQL','ManyToOne y FK hacia CENTRO_ACOPIO','Conserva la referencia entre registros.'),
], [1.05,3.12,2.58])
p('El frontend concentra las llamadas en OperacionesService, incluyendo las operaciones HTTP de recepciones; el componente usa sus métodos y no llama a HttpClient directamente. El servicio del backend busca el centro por ID, verifica que esté activo y construye la respuesta con la relación.',before=7,after=5)
callout('Distinción de modelos','La entrada necesita el ID para guardar la relación. La salida además incluye el nombre que la lista presenta al usuario. Así la interfaz no tiene que reconstruir esa relación a partir de texto libre.')
image_box(1,'Modelos y servicios de recepción','Muestra en VS Code los modelos de entrada y salida y los métodos HTTP de OperacionesService. Mantén visibles reloj y perfil.',2.32)

# PAGE 3 — list and parent filter.
page_break(); kicker('3.2 · Evidencia técnica'); heading('Lista con información relacionada',1)
p('La lista consulta las recepciones guardadas y presenta el nombre del centro de acopio asociado en cada fila. Para editar, también conserva idCentroAcopio como valor numérico en la respuesta.')
image_box(2,'Listado con Centro de Acopio','Muestra recepciones reales obtenidas desde el backend y la etiqueta del centro relacionado en cada fila.',2.05)
callout('Filtro por Centro de Acopio implementado','El listado permite seleccionar un centro. El frontend envía su identificador como parámetro idCentroAcopio; el endpoint, servicio y repositorio devuelven únicamente las recepciones asociadas. La petición GET con respuesta 200 se evidencia en Network.',fill='F0F7F2',title_color=BLUE)
image_box(3,'Filtro enviado en Network','La petición real GET /api/v1/mayorista/recepciones?idCentroAcopio=1 muestra el parámetro enviado al seleccionar Acopio Maria · La Rinconada.',1.92)

# PAGE 4 — dropdown and preselection.
page_break(); kicker('3.3 · Evidencia técnica'); heading('Lista desplegable y formulario',1)
p('El formulario carga centros desde el servicio HTTP de Centros de Acopio. El selector muestra nombre y zona, y guarda idCentroAcopio como número mediante ngValue. Al editar una recepción, el componente carga el mismo identificador en el control y Angular preselecciona la opción correspondiente.')
data_table(['Acción','Comportamiento'],[
('Nuevo registro','La opción inicial solicita elegir un centro de la lista.'),
('Edición','El formulario toma registro.idCentroAcopio y presenta el centro ya seleccionado.'),
('Tipo del valor','ngValue conserva el identificador numérico, no una cadena de texto.'),
], [1.35,5.4])
image_box(4,'Selector poblado desde el backend','Muestra el formulario real con centros cargados y la lista desplegable abierta o seleccionada.',2.25)
image_box(5,'Edición con Centro preseleccionado','Edita una recepción existente y deja visible el centro que ya estaba guardado.',2.05)

# PAGE 5 — dependency validation.
page_break(); kicker('3.4 · Evidencia técnica'); heading('Validación de dependencias',1)
p('La validación se realiza en dos niveles. En el formulario, idCentroAcopio inicia en cero y Validators.required junto con Validators.min(1) impiden guardar sin elección. En el backend, el request exige un ID positivo y RecepcionMayoristaService comprueba que el centro exista y siga activo antes de guardar.')
callout('Respuesta del backend','Si el centro enviado ya no existe o está inactivo, el servicio devuelve el error de recurso no encontrado. La pantalla informa que el centro dejó de estar disponible, reinicia el selector y vuelve a consultar centros activos.')
image_box(6,'Formulario sin Centro seleccionado','Intenta guardar sin seleccionar un Centro de Acopio y muestra el mensaje del formulario.',2.05)
image_box(7,'Centro inexistente o inactivo','Muestra el mensaje específico devuelto por la interfaz ante el error real del backend. No uses datos simulados.',2.05)

# PAGE 6 — CRUD connected to real backend.
page_break(); kicker('3.5 · Evidencia técnica'); heading('CRUD conectado al backend real',1)
p('El módulo de recepciones expone las operaciones de creación, consulta, edición y eliminación. El frontend mantiene la lista sincronizada con la respuesta del backend después de guardar o eliminar.')
data_table(['Operación','Petición del módulo','Resultado que debe evidenciarse'],[
('Listar','GET /api/v1/mayorista/recepciones','Registros que vienen del servidor.'),
('Crear','POST /api/v1/mayorista/recepciones','Respuesta persistida y fila nueva.'),
('Editar','PUT /api/v1/mayorista/recepciones/{id}','Cambio reflejado en la respuesta y la lista.'),
('Eliminar','DELETE /api/v1/mayorista/recepciones/{id}','Confirmación del servidor y fila retirada.'),
], [0.85,3.12,2.78])
image_box(8,'Crear una recepción','Registra una recepción real y muestra la confirmación y el registro persistido.',2.05)
image_box(9,'Editar y eliminar una recepción','Muestra la actualización persistida y, en otra captura si hace falta, la eliminación confirmada.',1.9)

# PAGE 7 — real finding, analysis and close.
page_break(); kicker('4 · Análisis'); heading('Error o hallazgo técnico',1)
p('Hallazgo real: las recepciones históricas guardaban el nombre del acopiador como texto, sin una clave foránea hacia CentroAcopio. Esto permitía que el texto quedara desactualizado o que no hubiera una relación que el sistema pudiera comprobar. Para corregirlo, se agregó la referencia estructural al centro y se ejecutó una migración que asoció las dos filas antiguas con coincidencia inequívoca.')
p('La revisión posterior confirmó dos recepciones con centro asociado y ninguna fila sin vínculo. También se añadió una prueba de integración para crear y actualizar recepciones, rechazar centros ausentes o inactivos y comprobar la lectura relacionada.',after=5)
image_box(10,'Hallazgo en migración y validación','Muestra el cambio de modelo/migración o la prueba del backend que diagnosticó el problema. Mantén reloj y perfil visibles.',2.0)
kicker('5 · Cierre'); heading('Reflexión técnica breve',1)
p('El formulario valida que se elija un centro para avisar el problema antes de enviar la petición y evitar intentos incompletos. El backend repite la validación porque también debe proteger los datos cuando la petición viene de otro cliente. Si solo existiera validación en el formulario, una petición directa podría enviar un ID inexistente. Si solo se validara en el backend, los datos estarían protegidos, pero el usuario recibiría el rechazo después de enviar. La clave foránea añade una protección en la base de datos. Las tres capas ayudan a mantener una relación válida y comprensible.',after=5)
callout('Actividad completada','El filtro por Centro de Acopio quedó implementado de extremo a extremo. La captura de Network muestra la petición real con idCentroAcopio=1 y respuesta 200; las demás capturas documentan el formulario, la validación y las operaciones CRUD realizadas en la aplicación.',fill='F0F7F2',title_color=BLUE)

# PAGE 8 — last page is the required feedback annex.
page_break(); q=p(align=WD_ALIGN_PARAGRAPH.RIGHT,after=2); run(q,'UNIVERSIDAD PERUANA UNIÓN',bold=True,color=NAVY,size=8.5)
q=p(align=WD_ALIGN_PARAGRAPH.RIGHT,after=8); run(q,'Escuela Profesional de Ingeniería de Sistemas',color=BLUE,size=8)
kicker('Última página del informe'); heading('Anexo · Feedback de la sesión',1)
p('Respuestas editables para revisión personal antes de exportar la entrega final.',after=7)
feedback=[
('1. ¿Cuál es el aprendizaje más importante que te llevas de la clase de hoy?','Aprendí que una tabla dependiente debe enviar el ID de la tabla padre y que la respuesta puede incluir su nombre para mostrar la relación.'),
('2. ¿Qué punto de la clase te resultó más confuso o te dejó con dudas?','Al inicio me costó diferenciar el ID que se guarda y el nombre que se muestra; todavía tengo dudas sobre el filtro que se envía al backend.'),
('3. ¿Tienes alguna pregunta que te gustaría que sea respondida la siguiente clase?','¿Cómo se implementa el filtro por tabla padre para que el backend reciba el parámetro y devuelva solo las filas relacionadas?'),
('4. Sobre tu nivel de comprensión de la clase, marca una opción:','[ ] ¡Entendido!  [X] Más o menos, entendí la idea general pero tengo dudas.  [ ] Necesito ayuda.'),
('5. ¿Cómo puedo ayudarte a comprender mejor el tema?','Con un ejemplo paso a paso que conecte formulario, modelos, servicio HTTP, endpoint, repositorio y respuesta.'),
('6. Pensando en tu participación y esfuerzo, marca una opción:','[ ] Muy comprometido/a.  [X] Comprometido/a.  [ ] Poco comprometido/a.'),
('7. Mi satisfacción con la clase fue (del 1 al 10):','Marca tu calificación:   1   2   3   4   5   6   7   8   9   10'),
]
for question,answer in feedback:
    q=p(after=1,keep=True); run(q,question,bold=True,color=NAVY,size=8.8)
    q=p(after=5,keep=True); run(q,answer,color=BODY,size=8.5)

doc.core_properties.title='Informe de evidencia S08 del módulo Mayorista'
doc.core_properties.subject='CRUD dependiente de Recepción Mayorista y Centro de Acopio'
doc.core_properties.author='Faijo Calisaya Helio Paul'
doc.core_properties.keywords='S08, LP2, SITRA-ORO, tablas dependientes, Centro de Acopio'
doc.save(OUT)
print(OUT)
