from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "S07_Mayorista_Plantilla_Capturas.docx"
OUT = ROOT / "docs" / "S07_Acopiador_Plantilla_Capturas.docx"


def paragraphs_in(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs_in(cell)


def replace_text(paragraph, mapping):
    for run in paragraph.runs:
        for old, new in mapping.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def blank_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main():
    doc = Document(SOURCE)
    replacements = {
        "SITRA-ORO · Aplicación al módulo Mayorista": "SITRA-ORO · Aplicación al módulo Acopiador",
        "Documentar la arquitectura SPA del frontend SITRA-ORO y comprobar la navegación y las operaciones del módulo Mayorista conectado al backend.": "Documentar la arquitectura SPA del frontend SITRA-ORO y comprobar la navegación y el CRUD independiente elegido para el módulo Acopiador, conectado al backend real.",
        "la ubicación de las pantallas del módulo Mayorista.": "la ubicación de las pantallas del módulo Acopiador.",
        "carpeta Mayorista.": "carpeta Acopiador.",
        "servicios propios del módulo Mayorista.": "servicios propios del módulo Acopiador.",
        "Navega a una ruta Mayorista y conserva el layout común visible.": "Navega a una ruta del módulo Acopiador y conserva el layout común visible.",
        "Muestra en VS Code el servicio Mayorista y traceIdInterceptor con crypto.randomUUID().": "Muestra en VS Code el servicio HTTP de tu CRUD independiente y el interceptor con crypto.randomUUID().",
        "OperacionesService concentra las llamadas del módulo y el interceptor añade un identificador de seguimiento a cada solicitud.": "El servicio HTTP concentra las llamadas de la entidad elegida y el interceptor añade un identificador de seguimiento a cada solicitud.",
        "En la pantalla de recepción apareció un mensaje genérico al fallar la consulta de registros.": "Describe brevemente el error o hallazgo real que encontraste al implementar o probar tu proyecto.",
        "CAPTURA 10. Hallazgo en la carga de recepciones": "CAPTURA 10. Evidencia del error o hallazgo",
        "El componente agrupaba varios errores bajo un mismo aviso, por lo que no orientaba sobre la causa. Se ajustó para distinguir sesión vencida, falta de permiso, ruta no disponible y otros fallos de conexión.": "Explica qué ocurrió, cómo lo comprobaste y qué hiciste para entenderlo o corregirlo.",
        "¿Por qué separar el servicio de la funcionalidad del componente facilita un cambio futuro en la URL o en la forma de consumir el backend? Responde en 5 a 8 líneas y menciona un servicio real de tu proyecto.": "¿Por qué separar el servicio HTTP de tu entidad del componente facilita un cambio futuro en la URL o en la forma de consumir el backend? Responde en 5 a 8 líneas.",
        "Faijo Calisaya Helio Paul": "____________________________",
        "Equipo 05": "Equipo: __________",
        "S07 · Creación y Arquitectura de la SPA": "S07 · Creación y Arquitectura de la SPA",
        "Frontend del módulo Mayorista: arquitectura SPA, navegación y conexión HTTP.": "Frontend del módulo Acopiador: ______________________________",
        "https://github.com/helionelio900-hub/SysProyec_Acopus": "____________________________________________",
        "Nombre": "Nombre",
    }
    for paragraph in paragraphs_in(doc):
        replace_text(paragraph, replacements)

    # Leave the identity fields for the teammate who will complete the report.
    for table in doc.tables:
        if len(table.rows) == 5 and len(table.columns) == 2:
            for row in table.rows:
                label = row.cells[0].text.strip()
                if label == "Equipo":
                    blank_paragraph(row.cells[1].paragraphs[0], "____________________________")
                elif label == "Sesión":
                    pass
                elif label == "Nombre":
                    blank_paragraph(row.cells[1].paragraphs[0], "____________________________")
                elif label == "Aporte":
                    blank_paragraph(row.cells[1].paragraphs[0], "____________________________")
                elif label == "GitHub":
                    blank_paragraph(row.cells[1].paragraphs[0], "____________________________")

    # Remove answers that belonged to the Mayorista report; the teammate supplies their own.
    answers = {
        "Separar OperacionesService del componente permite que el componente se concentre en la pantalla y que el servicio gestione las peticiones HTTP. Si cambia una ruta o la forma de consumir el backend, el ajuste se puede centralizar en el servicio, sin repetir cambios en cada pantalla. También mantiene tipadas las respuestas y facilita probar por separado la vista y la comunicación con el servidor.": "________________________________________________________________________________\n________________________________________________________________________________\n________________________________________________________________________________\n________________________________________________________________________________\n________________________________________________________________________________\n________________________________________________________________________________\n________________________________________________________________________________",
        "Aprendí a organizar una aplicación Angular en core, shared y features, y a conectar las pantallas con servicios e interceptores HTTP.": "Respuesta: __________________________________________________________________",
        "Me costó comprender al inicio cómo las rutas hijas se muestran dentro del layout sin reemplazar el encabezado y el menú.": "Respuesta: __________________________________________________________________",
        "¿Cómo se renueva un token JWT cuando varias solicitudes reciben un error 401 al mismo tiempo?": "Respuesta: __________________________________________________________________",
        "Con ejemplos prácticos de Signals y del manejo de estados de carga.": "Respuesta: __________________________________________________________________",
    }
    for paragraph in paragraphs_in(doc):
        if paragraph.text in answers:
            blank_paragraph(paragraph, answers[paragraph.text])
        elif paragraph.text == "[X] Entendido; podría explicarlo.":
            blank_paragraph(paragraph, "[ ] Entendido; podría explicarlo.")
        elif paragraph.text == "[X] Muy comprometido/a.":
            blank_paragraph(paragraph, "[ ] Muy comprometido/a.")
        elif paragraph.text.endswith("[X] 10"):
            blank_paragraph(paragraph, "1     2     3     4     5     6     7     8     9     10")

    for section in doc.sections:
        for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
            replace_text(paragraph, {"SITRA-ORO · Evidencia S07": "SITRA-ORO · Evidencia S07", "Faijo Calisaya Helio Paul · Equipo 05": "Acopiador · Equipo ______"})

    doc.core_properties.title = "Informe de evidencia S07 del módulo Acopiador"
    doc.core_properties.subject = "Plantilla para insertar capturas de evidencia"
    doc.core_properties.author = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
